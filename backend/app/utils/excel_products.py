from io import BytesIO
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.models.product import Product

SHEET_NAME = "Импорт товаров"
INSTRUCTIONS_SHEET_NAME = "Инструкция"

EXPORT_HEADERS = [
    "Артикул (SKU)", "Название", "Категория", "Подкатегория", "Бренд",
    "Цена", "Себестоимость", "Штрихкод", "Остаток", "Описание", "Детали",
]

IMPORT_HEADERS = [
    "Артикул (SKU)", "Название", "Категория", "Подкатегория", "Бренд",
    "Цена", "Себестоимость", "Штрихкод", "Описание", "Детали",
]

MARKER_ROW = [
    "необязательно", "ОБЯЗАТЕЛЬНО", "ОБЯЗАТЕЛЬНО", "необязательно", "необязательно",
    "ОБЯЗАТЕЛЬНО", "ОБЯЗАТЕЛЬНО", "необязательно", "необязательно", "необязательно",
]

_MARKER_WORDS = {"обязательно", "необязательно"}


def _autosize(ws: Worksheet, min_width: int = 10, max_width: int = 42) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=min_width)
        letter = col_cells[0].column_letter
        ws.column_dimensions[letter].width = min(max(length + 2, min_width), max_width)


def build_export_workbook(rows: list[dict]) -> BytesIO:
    """rows: list of dicts with keys matching EXPORT_HEADERS order (see repositories/product.py)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([
            row.get("sku") or "",
            row.get("name") or "",
            row.get("category") or "",
            row.get("subcategory") or "",
            row.get("brand") or "",
            row.get("price"),
            row.get("cost_price"),
            row.get("barcode") or "",
            row.get("available_quantity") or 0,
            row.get("description") or "",
            row.get("detail") or "",
        ])

    _autosize(ws)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_import_template_workbook(
    sample_category: Optional[str] = None,
    sample_subcategory: Optional[str] = None,
    sample_brand: Optional[str] = None,
) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    ws.append(IMPORT_HEADERS)
    ws.append(MARKER_ROW)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for cell in ws[2]:
        cell.font = Font(italic=True, color="9CA3AF")

    category = sample_category or "Электроника"
    subcategory = sample_subcategory or ""
    brand = sample_brand or "Generic"

    ws.append(["", "Ноутбук Example 15\"", category, subcategory, brand, 55000, 42000, "", "Пример: без артикула — будет создан новый товар", ""])
    ws.append(["EX-002", "Смартфон Example X", category, subcategory, brand, 25000, 18000, "", "Пример: с артикулом — повторный импорт этого файла обновит товар, а не создаст дубликат", ""])

    _autosize(ws)

    instructions = wb.create_sheet(INSTRUCTIONS_SHEET_NAME)
    instructions.append(["Колонка", "Обязательна", "Пояснение"])
    for cell in instructions[1]:
        cell.font = Font(bold=True)
    for row in [
        ("Артикул (SKU)", "Нет", "Если указан и совпадает с артикулом существующего товара — товар будет обновлён. Если не найден — создан новый. Если пусто — артикул будет сгенерирован автоматически."),
        ("Название", "Да", "Название товара."),
        ("Категория", "Да", "Точное название категории, как в системе (регистр не важен). Если не найдено — строка считается ошибочной и не импортируется."),
        ("Подкатегория", "Нет", "Точное название подкатегории, как в системе. Должна принадлежать указанной категории."),
        ("Бренд", "Нет", "Точное название бренда, как в системе."),
        ("Цена", "Да", "Цена продажи, число."),
        ("Себестоимость", "Да", "Закупочная цена, число."),
        ("Штрихкод", "Нет", "EAN-13 — 13 цифр с корректной контрольной суммой."),
        ("Описание", "Нет", "Краткое описание товара."),
        ("Детали", "Нет", "Дополнительная информация о товаре."),
    ]:
        instructions.append(row)
    _autosize(instructions)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _is_marker_row(values: tuple) -> bool:
    non_empty = [str(v).strip().lower() for v in values if v is not None and str(v).strip() != ""]
    if not non_empty:
        return False
    return all(word in _MARKER_WORDS for word in non_empty)


def parse_import_rows(content: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]

    rows_iter = ws.iter_rows(values_only=True)
    header_values = next(rows_iter, None)
    if not header_values:
        return []
    header = [str(h).strip() if h is not None else "" for h in header_values]

    rows: list[dict] = []
    excel_row_number = 1
    for values in rows_iter:
        excel_row_number += 1
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        if _is_marker_row(values):
            continue
        row = {header[i]: values[i] for i in range(min(len(header), len(values)))}
        row["_row_number"] = excel_row_number
        rows.append(row)

    wb.close()
    return rows
