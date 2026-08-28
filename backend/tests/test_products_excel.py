import io
import uuid

import pytest
from openpyxl import Workbook, load_workbook

pytestmark = pytest.mark.asyncio(loop_scope="session")

EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _uniq(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:8]}"


def _build_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Импорт товаров"
    ws.append([
        "Артикул (SKU)", "Название", "Категория", "Подкатегория", "Бренд",
        "Цена", "Себестоимость", "Штрихкод", "Описание", "Детали",
    ])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


async def _create_product(admin_client, **overrides):
    payload = {
        "name": _uniq("Product"),
        "description": None,
        "detail": None,
        "cost_price": 10,
        "price": 20,
        "category_id": None,
        "subcategory_id": None,
        "brand_id": None,
        "sku": _uniq("SKU"),
    }
    payload.update(overrides)
    resp = await admin_client.post("/products/", json=payload)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def test_import_template_downloads_valid_workbook(admin_client):
    resp = await admin_client.get("/products/import-template")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")

    wb = load_workbook(io.BytesIO(resp.content))
    assert "Импорт товаров" in wb.sheetnames
    assert "Инструкция" in wb.sheetnames

    header = [c.value for c in next(wb["Импорт товаров"].iter_rows(min_row=1, max_row=1))]
    assert "Название" in header
    assert "Категория" in header


async def test_import_preview_reports_row_errors_without_saving(admin_client):
    category_name = _uniq("Cat")
    resp = await admin_client.post("/categories/", json={"name": category_name})
    assert resp.status_code == 201

    content = _build_xlsx([
        ["", "Валидный товар", category_name, "", "", 100, 50, "", "", ""],
        ["", "Без категории", "Несуществующая категория XYZ", "", "", 100, 50, "", "", ""],
        ["", "Без цены", category_name, "", "", "", 50, "", "", ""],
    ])

    files = {"file": ("import.xlsx", content, EXCEL_CONTENT_TYPE)}
    resp = await admin_client.post("/products/import-excel/preview", files=files)
    assert resp.status_code == 200
    data = resp.json()

    assert data["summary"] == {"to_create": 1, "to_update": 0, "errors": 2}
    rows = data["rows"]
    assert rows[0]["action"] == "create"
    assert rows[1]["action"] == "error"
    assert any("не найдена" in e for e in rows[1]["errors"])
    assert rows[2]["action"] == "error"

    # preview must not persist anything
    resp = await admin_client.get("/products/", params={"name": "Валидный товар"})
    assert resp.json() == []


async def test_import_excel_creates_then_updates_on_matching_sku(admin_client):
    category_name = _uniq("Cat")
    resp = await admin_client.post("/categories/", json={"name": category_name})
    assert resp.status_code == 201

    sku = _uniq("SKU")
    content = _build_xlsx([
        [sku, "Новый товар", category_name, "", "", 150, 90, "", "Описание", ""],
    ])
    files = {"file": ("import.xlsx", content, EXCEL_CONTENT_TYPE)}
    resp = await admin_client.post("/products/import-excel", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["created"] == 1
    assert data["updated"] == 0
    assert data["errors"] == []

    content2 = _build_xlsx([
        [sku, "Обновлённый товар", category_name, "", "", 200, 120, "", "Другое описание", ""],
    ])
    files2 = {"file": ("import.xlsx", content2, EXCEL_CONTENT_TYPE)}
    resp = await admin_client.post("/products/import-excel", files=files2)
    assert resp.status_code == 200
    data2 = resp.json()
    assert data2["created"] == 0
    assert data2["updated"] == 1

    resp = await admin_client.get("/products/", params={"sku": sku})
    products = resp.json()
    assert len(products) == 1
    assert products[0]["name"] == "Обновлённый товар"
    assert products[0]["price"] == 200


async def test_import_excel_rejects_duplicate_sku_within_file(admin_client):
    category_name = _uniq("Cat")
    resp = await admin_client.post("/categories/", json={"name": category_name})
    assert resp.status_code == 201

    sku = _uniq("SKU")
    content = _build_xlsx([
        [sku, "Первая строка", category_name, "", "", 100, 50, "", "", ""],
        [sku, "Вторая строка с тем же SKU", category_name, "", "", 100, 50, "", "", ""],
    ])
    files = {"file": ("import.xlsx", content, EXCEL_CONTENT_TYPE)}
    resp = await admin_client.post("/products/import-excel/preview", files=files)
    data = resp.json()
    assert data["rows"][0]["action"] == "create"
    assert data["rows"][1]["action"] == "error"
    assert any("дублирующийся" in e for e in data["rows"][1]["errors"])


async def test_export_excel_respects_category_filter(admin_client):
    category_name = _uniq("ExportCat")
    resp = await admin_client.post("/categories/", json={"name": category_name})
    category_id = resp.json()["id"]
    other_category_name = _uniq("OtherCat")
    resp = await admin_client.post("/categories/", json={"name": other_category_name})
    other_category_id = resp.json()["id"]

    in_category = await _create_product(admin_client, category_id=category_id)
    outside_category = await _create_product(admin_client, category_id=other_category_id)

    resp = await admin_client.get("/products/export-excel", params={"category_id": category_id})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")

    wb = load_workbook(io.BytesIO(resp.content))
    skus = [row[0] for row in wb.active.iter_rows(min_row=2, values_only=True)]
    assert in_category["sku"] in skus
    assert outside_category["sku"] not in skus
